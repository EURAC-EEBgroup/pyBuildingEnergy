from pathlib import Path
import os
import sys
import zipfile


from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import MSO_AUTO_SIZE


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(os.environ.get("PBE_PPTX_TEMPLATE", ROOT / "EEB_conference_agenticAI.pptx"))
ISO_SCREENSHOT = Path(os.environ.get("PBE_ISO_SCREENSHOT", ROOT / "iso_screenshot.png"))
OUTPUT = ROOT / "pyBuildingEnergy_overview_EEB_layout_EN.pptx"

RED = RGBColor(226, 28, 19)
RED_DARK = RGBColor(179, 22, 17)
INK = RGBColor(54, 58, 63)
MUTED = RGBColor(112, 117, 122)
LIGHT = RGBColor(245, 245, 243)
PALE_RED = RGBColor(253, 238, 236)
PALE_BLUE = RGBColor(234, 241, 247)
BLUE = RGBColor(22, 71, 122)
GREEN = RGBColor(35, 126, 93)
WHITE = RGBColor(255, 255, 255)
FONT = "Calibri"
SERIF = "Calibri"


def delete_all_slides(prs):
    ids = list(prs.slides._sldIdLst)
    for slide_id in ids:
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             font=FONT, align=PP_ALIGN.LEFT, margin=0.03, valign=MSO_ANCHOR.TOP,
             italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return box


def add_runs(slide, runs, x, y, w, h, size=20, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    for item in runs:
        r = p.add_run()
        r.text = item[0]
        r.font.name = item[4] if len(item) > 4 else FONT
        r.font.size = Pt(item[1] if len(item) > 1 else size)
        r.font.bold = item[2] if len(item) > 2 else False
        r.font.color.rgb = item[3] if len(item) > 3 else INK
    return box


def rect(slide, x, y, w, h, fill=LIGHT, line=None, radius=True):
    typ = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    if radius:
        try: sh.adjustments[0] = 0.08
        except Exception: pass
    return sh


def line(slide, x1, y1, x2, y2, color=RED, width=2):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x1), Inches(y1), Inches(x2-x1), Inches(y2-y1))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def title(slide, text, kicker=None, num=None):
    if kicker:
        add_text(slide, kicker.upper(), 0.58, 0.22, 7.8, 0.23, 9.5, RED, True)
    add_text(slide, text, 0.58, 0.52, 12.0, 0.62, 27, INK, True, SERIF)
    line(slide, 0.58, 1.19, 12.17, 1.215, RED, 2)
    if num is not None:
        add_text(slide, f"{num:02}", 12.27, 0.50, 0.44, 0.35, 11, MUTED, True, align=PP_ALIGN.RIGHT)


def footer(slide, source=""):
    if source:
        add_text(slide, source, 0.60, 7.14, 11.6, 0.16, 6.7, MUTED)
    add_text(slide, "pyBuildingEnergy · Eurac Research", 10.55, 7.10, 2.15, 0.20, 7, RED, True, align=PP_ALIGN.RIGHT)


def bullet_list(slide, items, x, y, w, h, size=17, color=INK, gap=9):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02); tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = FONT; p.font.size = Pt(size); p.font.color.rgb = color
        p.level = 0; p.space_after = Pt(gap)
        p.bullet = True
    return box


def label(slide, text, x, y, w, color=RED, fill=PALE_RED):
    rect(slide, x, y, w, 0.34, fill)
    add_text(slide, text, x+0.07, y+0.07, w-0.14, 0.18, 9, color, True, align=PP_ALIGN.CENTER)


def add_hyperlink(box, url):
    for p in box.text_frame.paragraphs:
        for r in p.runs:
            r.hyperlink.address = url


prs = Presentation(str(TEMPLATE))
cover_layout = prs.slides[0].slide_layout
content_layout = prs.slides[1].slide_layout
contact_layout = prs.slides[-1].slide_layout
delete_all_slides(prs)

# 1 — Cover
s = prs.slides.add_slide(cover_layout)
for sh in list(s.shapes):
    if sh.is_placeholder:
        sh.element.getparent().remove(sh.element)
add_text(s, "pyBuildingEnergy", 3.05, 2.18, 9.65, 0.95, 43, RED, True, SERIF, PP_ALIGN.RIGHT)
line(s, 3.03, 3.21, 12.70, 3.245, RED, 2)
add_text(s, "From EPB standards to an open-source calculation engine", 4.05, 3.47, 8.65, 0.52, 20, RED_DARK, False, SERIF, PP_ALIGN.RIGHT)
add_text(s, "Transparent, reproducible and embeddable energy simulation", 4.50, 4.10, 8.20, 0.42, 15, INK, False, FONT, PP_ALIGN.RIGHT)
add_text(s, "Eurac Research · Institute for Renewable Energy", 8.05, 6.10, 4.65, 0.27, 11, INK, True, align=PP_ALIGN.RIGHT)
add_text(s, "September 2026", 9.95, 6.50, 2.75, 0.25, 10.5, RED, False, align=PP_ALIGN.RIGHT)

# 2 — What it does
s = prs.slides.add_slide(content_layout)
for sh in list(s.shapes):
    if sh.is_placeholder: sh.element.getparent().remove(sh.element)
title(s, "A Python engine for energy and comfort", "What it does", 2)
cards = [
    ("01", "Hourly energy needs", "Heating and cooling needs, indoor temperatures and heat balances according to EN ISO 52016-1."),
    ("02", "Envelope to systems", "EN ISO 52010/PVGIS weather, DHW EN 12831-3, distribution, generation, storage and ventilation."),
    ("03", "Reproducible workflows", "Structured inputs, quality checks, retrofit scenarios, reports and exportable results for platforms and APIs."),
]
for i,(n,h,b) in enumerate(cards):
    x=0.62+i*4.18
    rect(s,x,1.58,3.83,3.60,WHITE,RGBColor(225,225,225))
    add_text(s,n,x+0.20,1.82,0.55,0.52,24,RED,True,SERIF)
    add_text(s,h,x+0.20,2.45,3.38,0.45,18,INK,True)
    add_text(s,b,x+0.20,3.08,3.38,1.55,13.5,INK)
label(s,"OPEN SOURCE",0.62,5.60,1.45)
label(s,"PYTHON",2.28,5.60,1.10,BLUE,PALE_BLUE)
label(s,"HOURLY",3.60,5.60,1.15,GREEN,RGBColor(232,246,240))
add_text(s,"Useful for: research · pre-design · retrofit · benchmarking · embedded calculation engines",0.62,6.22,11.9,0.42,15,INK,True)
footer(s,"Sources: pyBuildingEnergy documentation and repository (accessed 20 Sep 2026).")

# 3 — Positioning
s = prs.slides.add_slide(content_layout)
for sh in list(s.shapes):
    if sh.is_placeholder: sh.element.getparent().remove(sh.element)
title(s,"It does not replace every simulator — it fills a different space","Positioning",3)
cols=[0.62,4.73,8.84]
data=[
 ("pyBuildingEnergy","Standards + embedded","A transparent Python engine aligned with EPB standards; quick to embed in services, batch workflows and platforms.",RED,PALE_RED),
 ("EnergyPlus / OpenStudio","Detailed dynamic simulation","Greater depth for HVAC, controls and complex geometries; higher modelling and computational effort.",BLUE,PALE_BLUE),
 ("IDA ICE / TRNSYS","Specialist environments","Strong for dynamic systems, co-simulation and professional workflows; often proprietary or less Python-backend-native.",GREEN,RGBColor(232,246,240)),
]
for x,(h,k,b,c,f) in zip(cols,data):
    rect(s,x,1.55,3.83,3.85,f)
    add_text(s,h,x+0.20,1.83,3.42,0.50,18,c,True,SERIF)
    add_text(s,k,x+0.20,2.45,3.42,0.32,11,c,True)
    add_text(s,b,x+0.20,3.05,3.42,1.58,13.2,INK)
add_runs(s,[("Key choice: ",16,True,RED),("when regulatory transparency, automation and scalability matter, pyBuildingEnergy reduces friction between calculation and digital product.",16,False,INK)],0.72,5.86,11.78,0.72)
footer(s,"Qualitative comparison: modelling depth depends on the use case and model configuration.")

# 4 — EPBD / ISO
s = prs.slides.add_slide(content_layout)
for sh in list(s.shapes):
    if sh.is_placeholder: sh.element.getparent().remove(sh.element)
title(s,"At the intersection of the EPBD and EPB standards","European relevance",4)
rect(s,0.62,1.52,4.65,4.75,PALE_RED)
add_text(s,"EPBD 2024",0.90,1.84,3.95,0.44,20,RED,True,SERIF)
add_text(s,"Member States must describe their national methodologies with reference to the annexes of the main EPB standards, including EN ISO 52010 and 52016.",0.90,2.50,3.95,1.35,14.5,INK)
add_text(s,"pyBuildingEnergy makes this logic:",0.90,4.10,3.95,0.35,13.5,INK,True)
bullet_list(s,["readable and verifiable","reusable in new services","accessible to the community"],0.96,4.58,3.65,1.25,13.2,INK,6)
s.shapes.add_picture(str(ISO_SCREENSHOT),Inches(5.62),Inches(1.52),width=Inches(7.08))
rect(s,6.02,5.53,6.25,0.72,WHITE,RGBColor(220,220,220))
add_text(s,"One of 3 European open-source tools presented at the ISO workshop",6.20,5.74,5.88,0.28,13,RED,True,align=PP_ALIGN.CENTER)
footer(s,"Sources: EPBD 2024 / pyBuildingEnergy documentation; ISO workshop screenshot supplied by the author.")

# 5 — SDPlus approach
s = prs.slides.add_slide(content_layout)
for sh in list(s.shapes):
    if sh.is_placeholder: sh.element.getparent().remove(sh.element)
title(s,"A faster physics-based route to Indian HVAC design loads","SDPlus · Approach",5)
add_text(s,"ISO 52016-1 with an Indian design-day weather file, implemented in pyBuildingEnergy and embedded in Smarter Dharma's SD+ workflow.",0.65,1.45,12.05,0.62,14.5,INK)
approach_cards=[
 ("01","The limitation","CLTD relies on generic reference tables and tends to overestimate. RTS depends on predefined response factors and lacks validated Indian assemblies.",RED,PALE_RED),
 ("02","The adaptation","Use real layer-by-layer construction properties in the hourly ISO 52016-1 RC network. Repeat one 24-hour design day for a 744-hour warm-up, then evaluate the final day.",BLUE,PALE_BLUE),
 ("03","Why it matters","A physics-based method becomes practical for repeated, platform-based HVAC design. Inputs and outputs are adapted around pyBuildingEnergy — its core remains unchanged.",GREEN,RGBColor(232,246,240)),
]
for i,(n,h,b,c,f) in enumerate(approach_cards):
    x=0.65+i*4.16
    rect(s,x,2.30,3.80,3.55,f)
    add_text(s,n,x+0.20,2.55,0.48,0.36,18,c,True)
    add_text(s,h,x+0.20,3.02,3.35,0.40,17,c,True)
    add_text(s,b,x+0.20,3.63,3.35,1.63,12.6,INK)
rect(s,0.65,6.12,12.00,0.54,WHITE,RGBColor(220,220,220))
add_runs(s,[("Main message  ",12,True,RED),("Replace generic construction assumptions with local data — without giving up speed or auditability.",13,True,INK)],0.88,6.28,11.50,0.23)
footer(s,"Source: SDPlus_ISO52016_two_slide_summary_v3.pptx, integrated and edited for consistency.")

# 6 — SDPlus validation
s = prs.slides.add_slide(content_layout)
for sh in list(s.shapes):
    if sh.is_placeholder: sh.element.getparent().remove(sh.element)
title(s,"Closer to EnergyPlus, with a much shorter runtime","SDPlus · Validation",6)
add_text(s,"Two real buildings in two climate zones provide an encouraging first result — not a final statistical validation.",0.65,1.45,12.05,0.48,14.5,INK)
chart_data=ChartData(); chart_data.categories=["Building 1","Building 2"]
chart_data.add_series("CLTD",(9.48,11.43)); chart_data.add_series("ISO 52016-1",(-4.60,6.15))
chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(0.70),Inches(2.18),Inches(7.15),Inches(3.77),chart_data).chart
chart.has_legend=True; chart.legend.position=XL_LEGEND_POSITION.BOTTOM; chart.legend.include_in_layout=False
chart.has_title=True; chart.chart_title.text_frame.text="Peak-load deviation from EnergyPlus (%)"
for p in chart.chart_title.text_frame.paragraphs: p.font.name = FONT
chart.legend.font.name = FONT
chart.value_axis.minimum_scale=-8; chart.value_axis.maximum_scale=14; chart.value_axis.major_unit=5
chart.value_axis.has_major_gridlines=True
chart.category_axis.tick_labels.font.size=Pt(11); chart.category_axis.tick_labels.font.name=FONT
chart.value_axis.tick_labels.font.size=Pt(10); chart.value_axis.tick_labels.font.name=FONT
chart.series[0].format.fill.solid(); chart.series[0].format.fill.fore_color.rgb=RGBColor(165,170,175)
chart.series[1].format.fill.solid(); chart.series[1].format.fill.fore_color.rgb=RED
for series in chart.series:
    series.has_data_labels=True; series.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
    series.data_labels.font.size=Pt(10); series.data_labels.font.name=FONT
rect(s,8.18,2.18,4.42,1.27,PALE_RED)
add_text(s,"9–16×",8.48,2.40,3.82,0.50,27,RED,True,SERIF,PP_ALIGN.CENTER)
add_text(s,"faster than EnergyPlus",8.48,2.91,3.82,0.24,11.5,INK,True,align=PP_ALIGN.CENTER)
rect(s,8.18,3.70,4.42,2.25,PALE_BLUE)
add_text(s,"What the numbers say",8.46,3.96,3.84,0.30,12,BLUE,True)
add_text(s,"CLTD is high in both cases: +9.5% and +11.4%. ISO 52016-1 is closer: −4.6% and +6.2%.",8.46,4.38,3.84,0.70,12.5,INK)
add_text(s,"The ISO error changes direction, so the accurate claim is “closer” — not “always conservative”.",8.46,5.17,3.84,0.53,11.5,INK,False,italic=True)
add_text(s,"Positive values indicate overestimation; negative values indicate underestimation.",0.78,6.20,7.00,0.25,10,MUTED)
footer(s,"Source: SDPlus summary. More buildings and climates are needed before generalising the result.")

# 7 — Ecosystem
s = prs.slides.add_slide(content_layout)
for sh in list(s.shapes):
    if sh.is_placeholder: sh.element.getparent().remove(sh.element)
title(s,"One calculation core, multiple platforms and contexts","Adoption",7)
add_text(s,"pyBuildingEnergy works as an infrastructure component: the same engine can power different interfaces, scenarios and services.",0.64,1.44,11.95,0.48,14.5,INK)
items=[
 ("FUTURHIST","Calculation basis for a platform supporting renovation of historic buildings.",RED,PALE_RED),
 ("ReLIFE","Engine for the European project's digital retrofit and energy-assessment services.",BLUE,PALE_BLUE),
 ("Cynergy","Calculation basis for an industry-commissioned project and its web tool.",GREEN,RGBColor(232,246,240)),
 ("Cerplan","Reused in internal tools for planning and scenario analysis.",RGBColor(134,82,27),RGBColor(250,242,231)),
]
for i,(h,b,c,f) in enumerate(items):
    x=0.65+(i%2)*6.18; y=2.18+(i//2)*1.60
    rect(s,x,y,5.75,1.30,f)
    add_text(s,h,x+0.18,y+0.20,1.55,0.28,12,c,True)
    add_text(s,b,x+1.82,y+0.18,3.65,0.76,12.3,INK)
rect(s,0.65,5.68,11.93,0.82,WHITE,RGBColor(220,220,220))
add_runs(s,[("Potential next connection: ",13,True,RED),("climatedataforbuildings.eu as an alternative weather source, proposed by NMBU through its frontend URL/API.",13,False,INK)],0.88,5.92,11.40,0.33)
footer(s,"Sources: supplied project websites; Arnkell Jonas Petersen (NMBU) email cited as a proposal, not a completed integration.")

# 8 — Community
s = prs.slides.add_slide(content_layout)
for sh in list(s.shapes):
    if sh.is_placeholder: sh.element.getparent().remove(sh.element)
title(s,"Measurable adoption, distributed development","Community & impact",8)
rect(s,0.64,1.53,3.32,2.12,PALE_RED)
add_text(s,"13k+",0.88,1.87,2.84,0.74,36,RED,True,SERIF,PP_ALIGN.CENTER)
add_text(s,"total downloads from PyPI",0.88,2.68,2.84,0.33,13,INK,True,align=PP_ALIGN.CENTER)
add_text(s,"pepy.tech data · 20 Sep 2026",0.88,3.12,2.84,0.22,9,MUTED,False,align=PP_ALIGN.CENTER)
add_text(s,"Human contributions to the library",4.35,1.57,8.05,0.42,18,INK,True,SERIF)
contrib=[
 ("Daniele Antonucci","direction, architecture, EPB workflows and systems integration"),
 ("Ulrich Filippi Oberegger","solver, building physics, components and reporting"),
 ("Kristian Stenerud Skeie · SINTEF","AHU and EN 16798-5-1 ventilation, tests and diagnostics"),
 ("Arnkell Jonas Petersen · NMBU","in-memory weather I/O, tests and workflow robustness"),
 ("Olga Somova","packaging, project structure and simulation tests"),
]
for i,(name,desc) in enumerate(contrib):
    y=2.18+i*0.73
    add_text(s,name,4.38,y,3.35,0.26,11.5,RED,True)
    add_text(s,desc,7.72,y,4.65,0.42,11.2,INK)
add_text(s,"The value is not only the code: it is the shared verifiability of methods, inputs and results.",0.73,4.12,3.10,1.26,15,INK,True,SERIF,PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
footer(s,"Sources: public pepy.tech badge; git shortlog and repository history. Human-attributed commits only; AI content excluded.")

# 9 — Close
s = prs.slides.add_slide(contact_layout)
for sh in list(s.shapes):
    if sh.is_placeholder: sh.element.getparent().remove(sh.element)
add_text(s,"pyBuildingEnergy",0.62,1.00,5.60,0.72,31,RED,True,SERIF)
line(s,0.62,1.82,5.62,1.85,RED,2)
add_text(s,"An open-source infrastructure for turning EPB standards into digital energy services.",0.62,2.18,5.45,1.50,20,INK,True,SERIF)
bullet_list(s,["transparent for research and verification","embeddable in platforms and products","ready to grow through external contributions"],0.68,4.05,5.20,1.40,13.5,INK,8)
rect(s,6.68,0.80,6.05,5.98,PALE_RED)
add_text(s,"Where to find it",7.14,1.28,5.10,0.42,19,RED,True,SERIF)
links=[
 ("GitHub","https://github.com/EURAC-EEBgroup/pyBuildingEnergy"),
 ("Documentation","https://eurac-eebgroup.github.io/pybuildingenergy-docs/"),
 ("PyPI / downloads","https://pepy.tech/projects/pybuildingenergy"),
]
for i,(txt,url) in enumerate(links):
    b=add_text(s,txt,7.15,2.10+i*0.72,4.95,0.34,14,BLUE,True)
    add_hyperlink(b,url)
add_text(s,"Invitation",7.15,4.55,4.95,0.30,11,RED,True)
add_text(s,"Use it, validate it on new cases, connect it to new data sources and contribute to the repository.",7.15,4.97,4.88,0.95,15,INK)
add_text(s,"Eurac Research · Institute for Renewable Energy",7.15,6.22,4.90,0.28,10,RED,True)

prs.core_properties.title = "pyBuildingEnergy — overview"
prs.core_properties.subject = "General presentation of the pyBuildingEnergy library"
prs.core_properties.author = "Eurac Research"
prs.core_properties.comments = "Created from the EEB Conference layout; data verified on 20 September 2026."
prs.save(str(OUTPUT))

# The EEB template contains Arial declarations in its layouts, masters and theme.
# Replace them too so every editable text element in the resulting deck resolves
# to Calibri, including inherited footer and chart formatting.
tmp_output = OUTPUT.with_suffix(".tmp.pptx")
with zipfile.ZipFile(OUTPUT, "r") as src, zipfile.ZipFile(tmp_output, "w", zipfile.ZIP_DEFLATED) as dst:
    for info in src.infolist():
        payload = src.read(info.filename)
        if info.filename.endswith(".xml") or info.filename.endswith(".rels"):
            payload = payload.replace(b'Arial', b'Calibri')
        dst.writestr(info, payload)
tmp_output.replace(OUTPUT)
print(OUTPUT)
